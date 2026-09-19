"""Helpers to edit the SD81 Booster manuals with python-docx instead of raw
XML string surgery. Reduces each addition to a few function calls instead of
hand-written OOXML blocks, while keeping the exact same visual styles
(dark code blocks, bold subheadings, table rows) already used in the manual.

Still required after using these helpers, same as before:
- Validate the result (open it, or run the existing validate.py from the
  docx skill).
- Render to PDF and look at the affected pages before trusting the output.

Usage example (see also test_docx_helpers.py):

    from docx import Document
    from docx_helpers import find_exact_paragraph, add_bold_paragraph, \
        add_code_block, add_body_paragraph, clone_table_row

    doc = Document("SD81_Manual_ES.docx")

    anchor = find_exact_paragraph(doc, "LOAD *64C")
    p = add_bold_paragraph(anchor, "Activar el modo de 256 caracteres:")
    p = add_code_block(p, "LOAD *256C")
    p = add_body_paragraph(p, "Explicacion...")

    table = doc.tables[N]
    clone_table_row(table, table.rows[-1], ["65", "SEL_256CHARS", None, None,
                                             "Descripcion..."])

    doc.save("SD81_Manual_ES.docx")
"""
import copy

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

# Molds extracted from the manual's own styles (see
# reference_manual_docx.md in memory for where these come from).
CODE_FILL = "1E1E1E"
CODE_FONT_COLOR = RGBColor(0xA8, 0xD8, 0xA8)
CODE_FONT_NAME = "Courier New"
CODE_FONT_SIZE_PT = 10  # sz val="20" half-points


def find_exact_paragraph(doc, text):
    """Return the single top-level paragraph whose text equals `text`
    exactly. Raises if none or more than one match (e.g. the TOC caches
    the same heading text as the body -- use a more specific anchor, like
    an existing code-block line, to avoid that ambiguity)."""
    matches = [p for p in doc.paragraphs if p.text == text]
    if not matches:
        raise ValueError(f"No paragraph found with exact text: {text!r}")
    if len(matches) > 1:
        raise ValueError(
            f"Ambiguous: {len(matches)} paragraphs match {text!r} -- "
            "pick a more specific anchor (e.g. a code-block line, which "
            "usually doesn't also appear in the TOC)."
        )
    return matches[0]


def insert_paragraph_after(ref_paragraph):
    """Create a new, empty paragraph immediately after `ref_paragraph` and
    return it. Chain calls using the paragraph just returned to keep
    inserting further down (see module docstring example)."""
    new_p = OxmlElement("w:p")
    ref_paragraph._p.addnext(new_p)
    return Paragraph(new_p, ref_paragraph._parent)


def set_spacing(paragraph, before=None, after=None):
    """before/after in the same units as the manual's raw XML
    (twentieths of a point, e.g. spacing before="80" -> before=80)."""
    pf = paragraph.paragraph_format
    if before is not None:
        pf.space_before = Pt(before / 20)
    if after is not None:
        pf.space_after = Pt(after / 20)


def set_shading(paragraph, fill_hex):
    """CT_PPrBase is schema-ordered: w:shd must come before w:spacing/
    w:ind/etc, so it can't just be appended -- it has to go first."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.insert(0, shd)


def add_bold_paragraph(ref_paragraph, text):
    """Mold: body paragraph, bold run (e.g. 'Activar el modo de X:')."""
    p = insert_paragraph_after(ref_paragraph)
    set_spacing(p, before=60, after=100)
    run = p.add_run(text)
    run.bold = True
    return p


def add_code_block(ref_paragraph, text):
    """Mold: dark-background code line (LOAD *xxx examples)."""
    p = insert_paragraph_after(ref_paragraph)
    set_spacing(p, before=60, after=60)
    p.paragraph_format.left_indent = Pt(36)  # w:ind w:left="720" in the mold
    set_shading(p, CODE_FILL)
    run = p.add_run(text)
    run.font.name = CODE_FONT_NAME
    run.font.size = Pt(CODE_FONT_SIZE_PT)
    run.font.color.rgb = CODE_FONT_COLOR
    return p


def add_body_paragraph(ref_paragraph, text):
    """Mold: plain explanatory paragraph."""
    p = insert_paragraph_after(ref_paragraph)
    set_spacing(p, before=80, after=80)
    p.add_run(text)
    return p


def set_cell_text(cell, text):
    """Replace a table cell's text while KEEPING the formatting of its
    first run (color, font, size) -- e.g. the orange Courier New used for
    the address column in the POKE summary table. Extra runs/paragraphs
    beyond the first are dropped."""
    paragraphs = cell.paragraphs
    first_para = paragraphs[0]
    if first_para.runs:
        first_run = first_para.runs[0]
    else:
        first_run = first_para.add_run()
    first_run.text = text
    for r in first_para.runs[1:]:
        r._r.getparent().remove(r._r)
    for p in paragraphs[1:]:
        p._p.getparent().remove(p._p)


def add_heading_after(ref_element, parent, text, style="Ttulo2"):
    """Insert a new heading paragraph right after `ref_element` (either a
    Paragraph's `._p` or a Table's `._tbl` -- anything with `.addnext()`)
    and return it as a Paragraph. style is one of the manual's heading
    styles: 'Ttulo1' (chapter/appendix), 'Ttulo2' (section), 'Ttulo3'
    (sub-table heading). No TOC bookmark is added -- matches the pattern
    used by the most recently added appendices (I, J), which rely on the
    user refreshing the TOC in Word (F9) rather than a pre-existing
    _TocNNNN bookmark."""
    new_p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    pstyle = OxmlElement("w:pStyle")
    pstyle.set(qn("w:val"), style)
    ppr.append(pstyle)
    new_p.append(ppr)
    ref_element.addnext(new_p)
    p = Paragraph(new_p, parent)
    p.add_run(text)
    return p


def clone_table_after(ref_element, template_table, keep_rows=2):
    """Deep-copy `template_table` (borders/shading/column widths and all),
    insert it right after `ref_element`, trim it down to the first
    `keep_rows` rows (typically the header row plus one data row to use as
    the formatting template for clone_table_row), and return the new
    Table. Caller then does set_cell_text on the kept data row and/or
    clone_table_row for any further rows."""
    new_tbl_el = copy.deepcopy(template_table._tbl)
    trs = new_tbl_el.findall(qn("w:tr"))
    for tr in trs[keep_rows:]:
        new_tbl_el.remove(tr)
    ref_element.addnext(new_tbl_el)
    return Table(new_tbl_el, template_table._parent)


def clone_table_row(table, template_row, cell_texts):
    """Append a new row to `table`, cloned from `template_row` (same
    borders/shading/per-cell formatting), with each cell's text replaced
    by the corresponding entry in `cell_texts` (use None to leave a cell
    exactly as in the template, e.g. the '---' placeholder cells)."""
    new_tr = copy.deepcopy(template_row._tr)
    table._tbl.append(new_tr)
    new_row = _Row(new_tr, table)
    for cell, text in zip(new_row.cells, cell_texts):
        if text is not None:
            set_cell_text(cell, text)
    return new_row
